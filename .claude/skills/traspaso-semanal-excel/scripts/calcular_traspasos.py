#!/usr/bin/env python3
"""
Calcula las sugerencias de traspaso semanal entre tiendas de NUNU (AMIGO,
MADRID, RAMBLA, VALENCIA) a partir de un Excel con stock y ventas por
tienda, y escribe una hoja nueva con la misma tabla que se ve en la
pestaña "Revisar traspasos" de la app Repo_semanal.

Este script es un port fiel a Python de computeMovimientosParaFila (y sus
funciones auxiliares) en app.js — misma lógica, mismo orden de fases,
mismos criterios de desempate. Cualquier cambio de negocio en el
algoritmo debe hacerse primero en app.js / CRITERIOS_REPO_SEMANAL.md y
después reflejarse aquí; no al revés.

Uso:
    python calcular_traspasos.py entrada.xlsx
    python calcular_traspasos.py entrada.xlsx --sheet "Datos" --output salida.xlsx
    python calcular_traspasos.py entrada.xlsx --mapping mapping.json --config config.json

Por defecto escribe (o sobrescribe) la hoja "Traspasos" dentro del propio
archivo de entrada, dejando el resto de hojas intactas.
"""

import argparse
import json
import re
import sys
import unicodedata

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
except ImportError:
    import subprocess
    subprocess.run([sys.executable, '-m', 'pip', 'install', '--quiet', 'openpyxl'], check=True)
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter


STORES = ['AMIGO', 'MADRID', 'RAMBLA', 'VALENCIA']
STORE_LABELS = {'AMIGO': 'Amigó', 'MADRID': 'Madrid', 'RAMBLA': 'Rambla', 'VALENCIA': 'Valencia'}
STORE_ABBR = {'AMIGO': 'AMI', 'MADRID': 'MAD', 'RAMBLA': 'RAM', 'VALENCIA': 'VLC'}
MOVEMENT_ORDER = [(f, t) for f in STORES for t in STORES if t != f]

DEFAULT_CONFIG = {
    'top20ExcludePrefixes': ['ACC', 'GC'],
    'top20ExcludeKeywords': [
        'gift card', 'tarjeta regalo', 'envoltorio', 'papel de regalo',
        'bolsa de regalo', 'accesorio', 'gift wrap',
    ],
    'repoImportExcludeSkuPrefixes': ['GIFT'],
    'repoAccessorySkuPrefixes': ['ACC'],
    'repoAccessoryKeywords': ['accesorio'],
    'repoBufferMadrid': 2,
    'repoBufferRamblaValencia': 1,
    'repoAmigoProtectedMargin': 1,
    'repoAmigoTop20MinStock': 3,
    'whPrefix': 'WH-',
}


# ============ Detección de columnas de entrada ============
# Mismo criterio de "adivinar" columnas que usa la app al importar un CSV
# (ver storeGuessRegex / only1mSalesGuess en app.js): primero intenta un
# nombre canónico exacto (para el formato de ida y vuelta de la propia
# app: SKU, Producto, Stock_AMIGO, Ventas_AMIGO...), y si no lo encuentra
# cae a una búsqueda difusa por palabras clave.

def strip_accents(s):
    return ''.join(c for c in unicodedata.normalize('NFKD', s) if not unicodedata.combining(c))


def canon(s):
    return re.sub(r'[^a-z0-9]+', '', strip_accents(str(s or '')).lower())


def only_1m_sales_pattern(base):
    # Igual que only1mSalesGuess en app.js: evita columnas de ventas a
    # varios meses ("ventas 3 meses", "sales 12m"...) para no confundirlas
    # con la venta de la última ventana (1 mes) que usa el algoritmo.
    return re.compile(rf'(?!.*([2-9]|1[0-2])\s*m(es(es)?)?\b)(?:{base})', re.IGNORECASE)


def store_guess_pattern(store, kind):
    abbr = re.escape(STORE_ABBR[store].lower())
    label = re.escape(STORE_LABELS[store].lower())
    kind_pat = r'(stock|available|existenc)' if kind == 'stock' else r'(sales|ventas|vendid)'
    base = rf'({abbr}|{label}).*{kind_pat}|{kind_pat}.*({abbr}|{label})'
    return only_1m_sales_pattern(base) if kind == 'sales' else re.compile(base, re.IGNORECASE)


RE_SKU = re.compile(r'^sku$|codigo', re.IGNORECASE)
RE_NOMBRE = re.compile(r'product|nombre|descrip|name', re.IGNORECASE)
RE_ONLINE = only_1m_sales_pattern('online')


def detect_columns(headers, explicit_mapping=None):
    """headers: lista de encabezados (en orden). Devuelve dict field -> header,
    o None si no se pudo determinar (para campos opcionales)."""
    explicit_mapping = explicit_mapping or {}
    mapping = {}
    warnings = []

    canon_headers = {canon(h): h for h in headers}

    def resolve(field, canonical_names, regex, required):
        if field in explicit_mapping:
            header = explicit_mapping[field]
            if header not in headers:
                raise SystemExit(f'--mapping: la columna "{header}" indicada para "{field}" no existe en la hoja. Encabezados disponibles: {headers}')
            mapping[field] = header
            return
        for name in canonical_names:
            if name in canon_headers:
                mapping[field] = canon_headers[name]
                return
        for h in headers:
            if regex.search(str(h)):
                mapping[field] = h
                return
        if required:
            raise SystemExit(
                f'No se pudo identificar la columna "{field}" en el Excel. '
                f'Encabezados encontrados: {headers}. '
                f'Renombra la columna o indícala explícitamente con --mapping.'
            )
        mapping[field] = None
        warnings.append(field)

    resolve('sku', ['sku', 'codigo', 'referencia'], RE_SKU, required=True)
    resolve('nombre', ['nombre', 'producto', 'descripcion'], RE_NOMBRE, required=True)
    for s in STORES:
        resolve(f'stock_{s}', [f'stock{s.lower()}', f'stock{STORE_ABBR[s].lower()}'],
                store_guess_pattern(s, 'stock'), required=False)
        resolve(f'sales_{s}', [f'ventas{s.lower()}', f'ventas{STORE_ABBR[s].lower()}', f'sales{s.lower()}'],
                store_guess_pattern(s, 'sales'), required=False)
    resolve('sales_ONLINE', ['ventasonline', 'salesonline', 'online'], RE_ONLINE, required=False)

    return mapping, warnings


# ============ Lectura y agregación de filas (equivalente al import CSV) ============

def parse_number(v):
    if v is None or v == '':
        return None
    if isinstance(v, (int, float)):
        if isinstance(v, float) and v != v:  # NaN
            return None
        return int(round(v))
    try:
        return int(round(float(str(v).strip().replace(',', '.'))))
    except ValueError:
        return None


def normalize_sku(v):
    return str(v).strip().upper() if v is not None else ''


def is_wholesale_row(sku, nombre, cfg):
    prefix = (cfg.get('whPrefix') or 'WH-').upper()
    sku_up = (sku or '').upper()
    nombre_up = (nombre or '').upper().strip()
    return (bool(prefix) and sku_up.startswith(prefix)) or nombre_up.startswith('(WH)')


def is_excluded_from_import(sku, cfg):
    sku_up = (sku or '').upper()
    return any(p and sku_up.startswith(p.upper()) for p in cfg['repoImportExcludeSkuPrefixes'])


def is_accessory_row(row, cfg):
    sku = (row.get('sku') or '').upper()
    nombre = (row.get('nombre') or '').upper()
    if any(p and sku.startswith(p.upper()) for p in cfg['repoAccessorySkuPrefixes']):
        return True
    if any(k and k.upper() in nombre for k in cfg['repoAccessoryKeywords']):
        return True
    return False


def is_real_product_for_top20(row, cfg):
    sku = (row.get('sku') or '').upper()
    nombre = (row.get('nombre') or '').upper()
    if any(p and sku.startswith(p.upper()) for p in cfg['top20ExcludePrefixes']):
        return False
    if any(k and k.upper() in nombre for k in cfg['top20ExcludeKeywords']):
        return False
    return True


def build_rows(sheet_rows, headers, mapping, cfg):
    """sheet_rows: lista de dicts header->valor (una entrada por fila de datos)."""
    by_sku = {}
    rows = []
    stats = {'wholesale': 0, 'excluded': 0, 'negativos': [], 'filas_leidas': len(sheet_rows)}

    for r in sheet_rows:
        sku = normalize_sku(r.get(mapping['sku']))
        if not sku:
            continue
        nombre_raw = r.get(mapping['nombre']) if mapping['nombre'] else None
        nombre_raw = (str(nombre_raw).strip() if nombre_raw not in (None, '') else '') or sku
        if is_wholesale_row(sku, nombre_raw, cfg):
            stats['wholesale'] += 1
            continue
        if is_excluded_from_import(sku, cfg):
            stats['excluded'] += 1
            continue
        entry = by_sku.get(sku)
        if not entry:
            entry = {
                'sku': sku, 'nombre': nombre_raw,
                'stock': {s: 0 for s in STORES}, 'sales': {s: 0 for s in STORES},
                'top20': {s: False for s in STORES},
            }
            by_sku[sku] = entry
            rows.append(entry)
        for s in STORES:
            stock_col = mapping.get(f'stock_{s}')
            sales_col = mapping.get(f'sales_{s}')
            if stock_col:
                v = parse_number(r.get(stock_col))
                if v is not None:
                    entry['stock'][s] += v
                    if v < 0:
                        stats['negativos'].append((sku, f'Stock {STORE_ABBR[s]}'))
            if sales_col:
                v = parse_number(r.get(sales_col))
                if v is not None:
                    entry['sales'][s] += v
                    if v < 0:
                        stats['negativos'].append((sku, f'Ventas {STORE_ABBR[s]}'))
        online_col = mapping.get('sales_ONLINE')
        if online_col:
            v = parse_number(r.get(online_col))
            if v is not None:
                entry['sales']['AMIGO'] += v
                if v < 0:
                    stats['negativos'].append((sku, 'Ventas Online'))

    return rows, stats


# ============ Top20 ============

def compute_top20(rows, cfg):
    for s in STORES:
        eligible = [r for r in rows if is_real_product_for_top20(r, cfg) and r['sales'].get(s, 0) > 0]
        eligible.sort(key=lambda r: r['sales'].get(s, 0), reverse=True)
        top_skus = {r['sku'] for r in eligible[:20]}
        for r in rows:
            r['top20'][s] = r['sku'] in top_skus


# ============ Consolidación de cadenas intermedias ============
# A->B->C se convierte en A->C directo (evita traspasos redundantes que en
# realidad esconden un movimiento directo).

def consolidar_cadenas_intermedias(movimientos):
    movimientos = [dict(m) for m in movimientos]
    cambiado = True
    while cambiado:
        cambiado = False
        for intermedia in STORES:
            entrantes = [m for m in movimientos if m['to'] == intermedia and m['qty'] > 0]
            salientes = [m for m in movimientos if m['from'] == intermedia and m['qty'] > 0]
            for inc in entrantes:
                for out in salientes:
                    if inc['qty'] <= 0 or out['qty'] <= 0:
                        continue
                    qty = min(inc['qty'], out['qty'])
                    if qty <= 0:
                        continue
                    if inc['from'] == out['to']:
                        inc['qty'] -= qty
                        out['qty'] -= qty
                        cambiado = True
                        continue
                    inc['qty'] -= qty
                    out['qty'] -= qty
                    directo = next((m for m in movimientos if m['from'] == inc['from'] and m['to'] == out['to']), None)
                    if directo:
                        directo['qty'] += qty
                    else:
                        movimientos.append({'from': inc['from'], 'to': out['to'], 'qty': qty})
                    cambiado = True
        movimientos = [m for m in movimientos if m['qty'] > 0]
    return movimientos


# ============ Algoritmo principal: computeMovimientosParaFila ============
# Port 1:1 de app.js. Ver CRITERIOS_REPO_SEMANAL.md para la explicación de
# negocio de cada fase; los comentarios aquí solo marcan qué fase es cada
# bloque para poder comparar con el original línea a línea.

def compute_movimientos_para_fila(row, cfg):
    if is_accessory_row(row, cfg):
        return []

    stock = {s: row['stock'].get(s, 0) for s in STORES}
    giveable = {s: row['stock'].get(s, 0) for s in STORES}
    sales = {s: row['sales'].get(s, 0) for s in STORES}
    movimientos = []
    extreme = {s: bool(row['top20'].get(s)) for s in STORES}

    def full_need_target(store):
        if store == 'MADRID':
            return sales['MADRID'] + cfg['repoBufferMadrid']
        if store == 'AMIGO':
            return sales['AMIGO']
        return sales[store] + cfg['repoBufferRamblaValencia']

    def add_mov(frm, to, qty):
        if qty <= 0:
            return
        existing = next((m for m in movimientos if m['from'] == frm and m['to'] == to), None)
        if existing:
            existing['qty'] += qty
        else:
            movimientos.append({'from': frm, 'to': to, 'qty': qty})
        stock[frm] -= qty
        stock[to] += qty
        giveable[frm] -= qty

    normales = [s for s in STORES if not extreme[s]]
    escasez = sum(stock[s] for s in normales) < sum(sales[s] for s in normales)

    def floor_amigo_normal():
        return max(sales['AMIGO'] + (-1 if escasez else cfg['repoAmigoProtectedMargin']), 1)

    def floor_origen(store):
        return max(sales[store] + (-1 if escasez else 0), 1)

    destinos = [{'store': s, 'shortage': full_need_target(s) - stock[s]} for s in STORES]
    destinos = [d for d in destinos if d['shortage'] > 0]
    destinos.sort(key=lambda d: d['shortage'], reverse=True)

    # 1) Casos extremos (top20 del destino) primero.
    if extreme['AMIGO']:
        extremos = [s for s in STORES if extreme[s]]
        for _ in range(200):
            con_margen = sorted(
                ({'store': s, 'diff': stock[s] - sales[s]} for s in extremos),
                key=lambda x: x['diff'], reverse=True,
            )
            origen = con_margen[0]
            destino = con_margen[-1]
            if origen['store'] == destino['store'] or origen['diff'] <= destino['diff']:
                break
            if giveable[origen['store']] <= 1:
                break
            add_mov(origen['store'], destino['store'], 1)
    else:
        for dest in [d for d in destinos if extreme[d['store']]]:
            if dest['store'] == 'AMIGO':
                continue
            falta = full_need_target(dest['store']) - stock[dest['store']]
            if falta <= 0:
                continue
            floor_top20 = max(cfg['repoAmigoTop20MinStock'], 1)
            floor_amigo = min(floor_top20, floor_amigo_normal())
            disponible = giveable['AMIGO'] - floor_amigo
            if disponible > 0:
                add_mov('AMIGO', dest['store'], min(disponible, falta))

    # 2) Tiendas normales, en pasadas de prioridad decreciente.
    # a) Venta base (breakeven).
    for _ in range(200):
        necesitadas = [s for s in normales if sales[s] - stock[s] > 0]
        if not necesitadas:
            break
        necesitadas.sort(key=lambda s: stock[s] - sales[s])
        destino = necesitadas[0]
        origen = None
        if destino != 'AMIGO' and giveable['AMIGO'] - floor_amigo_normal() > 0:
            origen = 'AMIGO'
        else:
            candidatos = [s for s in normales if s != destino and s != 'AMIGO' and giveable[s] - floor_origen(s) > 0]
            candidatos.sort(key=lambda s: stock[s] - sales[s], reverse=True)
            if candidatos:
                origen = candidatos[0]
        if not origen:
            break
        add_mov(origen, destino, 1)

    if not escasez:
        # b) Presencia mínima.
        for destino in [s for s in normales if s != 'AMIGO' and stock[s] == 0]:
            disponible = giveable['AMIGO'] - floor_amigo_normal()
            if disponible > 0:
                add_mov('AMIGO', destino, 1)

        # c) Margen extra.
        margen_targets = [{'store': s, 'shortage': full_need_target(s) - stock[s]} for s in normales if sales[s] > 0]
        margen_targets = [d for d in margen_targets if d['shortage'] > 0]
        margen_targets.sort(key=lambda d: d['shortage'], reverse=True)
        for dest in margen_targets:
            falta = full_need_target(dest['store']) - stock[dest['store']]
            if falta <= 0:
                continue
            if dest['store'] != 'AMIGO':
                disponible = giveable['AMIGO'] - floor_amigo_normal()
                if disponible > 0:
                    qty = min(disponible, falta)
                    add_mov('AMIGO', dest['store'], qty)
                    falta -= qty
            if falta <= 0:
                continue
            if dest['store'] == 'MADRID':
                falta = max(0, (sales['MADRID'] + cfg['repoBufferRamblaValencia']) - stock['MADRID'])
                if falta <= 0:
                    continue
            candidatos = [{'store': s, 'disponible': giveable[s] - floor_origen(s)} for s in normales if s != 'AMIGO' and s != dest['store']]
            candidatos = [c for c in candidatos if c['disponible'] > 0]
            candidatos.sort(key=lambda c: c['disponible'], reverse=True)
            for c in candidatos:
                if falta <= 0:
                    break
                qty = min(c['disponible'], falta)
                add_mov(c['store'], dest['store'], qty)
                falta -= qty

        # d) Aprovechar sobrante (roturas reales).
        if is_real_product_for_top20(row, cfg):
            candidatos_sobrante = [s for s in normales if s != 'AMIGO' and sales[s] > 0 and row['stock'].get(s, 0) == 0]
            for _ in range(200):
                if giveable['AMIGO'] - floor_amigo_normal() <= 0:
                    break
                peores = sorted((s for s in candidatos_sobrante if stock[s] < stock['AMIGO']), key=lambda s: stock[s])
                if not peores:
                    break
                add_mov('AMIGO', peores[0], 1)

    # 3) Rescate de emergencia.
    rescate_targets = [s for s in STORES if sales[s] > 0 and stock[s] <= 1 and stock[s] < sales[s]]
    rescate_targets.sort(key=lambda s: sales[s] - stock[s], reverse=True)
    for destino in rescate_targets:
        candidatos = [s for s in STORES if s != destino and giveable[s] > 1]
        candidatos.sort(key=lambda s: stock[s] - sales[s], reverse=True)
        if candidatos:
            add_mov(candidatos[0], destino, 1)

    # Último recurso.
    if not movimientos:
        en_deficit = [s for s in STORES if stock[s] < sales[s]]
        if en_deficit:
            en_deficit.sort(key=lambda s: sales[s] - stock[s], reverse=True)
            destino = en_deficit[0]
            candidatos = [s for s in STORES if s != destino and giveable[s] > 1]
            candidatos.sort(key=lambda s: stock[s] - sales[s], reverse=True)
            if candidatos:
                add_mov(candidatos[0], destino, 1)

    return consolidar_cadenas_intermedias(movimientos)


def mov_qty_for_pair(movimientos, frm, to):
    for m in movimientos:
        if m['from'] == frm and m['to'] == to:
            return m['qty']
    return 0


# ============ Orden por nombre (equivalente a sortByNombre, es-ES base) ============

def sort_key_nombre(nombre):
    return canon(nombre or '')


# ============ Lectura del Excel de entrada ============

def read_input_sheet(path, sheet_name):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet_name] if sheet_name else wb[wb.sheetnames[0]]
    rows_iter = ws.iter_rows(values_only=True)
    try:
        headers = [str(h) if h is not None else '' for h in next(rows_iter)]
    except StopIteration:
        raise SystemExit(f'La hoja "{ws.title}" está vacía.')
    sheet_rows = []
    for values in rows_iter:
        if values is None or all(v is None for v in values):
            continue
        sheet_rows.append(dict(zip(headers, values)))
    return ws.title, headers, sheet_rows


# ============ Escritura de la hoja de salida ============

HEADER_FILL = PatternFill('solid', fgColor='1F2937')
HEADER_FONT = Font(color='FFFFFF', bold=True)
MOV_FILL = PatternFill('solid', fgColor='DCFCE7')
TOP20_FONT = Font(color='166534', bold=True)


def write_output_sheet(path, out_sheet_name, rows, cfg):
    wb = openpyxl.load_workbook(path, data_only=False)
    if out_sheet_name in wb.sheetnames:
        del wb[out_sheet_name]
    ws = wb.create_sheet(out_sheet_name)

    header = ['Producto', 'SKU']
    for s in STORES:
        header += [f'Top20 {STORE_LABELS[s]}']
    for s in STORES:
        header += [f'Stock {STORE_LABELS[s]}', f'Ventas {STORE_LABELS[s]}']
    for (f, t) in MOVEMENT_ORDER:
        header.append(f'{STORE_ABBR[f]}→{STORE_ABBR[t]}')
    ws.append(header)
    for col_idx in range(1, len(header) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal='center', wrap_text=True)

    ordered = sorted(rows, key=lambda r: sort_key_nombre(r['nombre']))
    for r in ordered:
        movimientos = r['movimientos']
        row_values = [r['nombre'], r['sku']]
        row_values += ['SI' if r['top20'][s] else 'NO' for s in STORES]
        for s in STORES:
            row_values += [r['stock'].get(s, 0), r['sales'].get(s, 0)]
        for (f, t) in MOVEMENT_ORDER:
            row_values.append(mov_qty_for_pair(movimientos, f, t))
        ws.append(row_values)

    n_data_cols = len(header)
    n_data_rows = len(ordered) + 1
    top20_col_start = 3
    mov_col_start = 3 + len(STORES) + 2 * len(STORES)
    for row_idx in range(2, n_data_rows + 1):
        for c in range(top20_col_start, top20_col_start + len(STORES)):
            cell = ws.cell(row=row_idx, column=c)
            if cell.value == 'SI':
                cell.font = TOP20_FONT
        for c in range(mov_col_start, mov_col_start + len(MOVEMENT_ORDER)):
            cell = ws.cell(row=row_idx, column=c)
            if isinstance(cell.value, (int, float)) and cell.value > 0:
                cell.fill = MOV_FILL

    ws.freeze_panes = 'C2'
    ws.auto_filter.ref = f'A1:{get_column_letter(n_data_cols)}{n_data_rows}'
    for c in range(1, n_data_cols + 1):
        letter = get_column_letter(c)
        header_len = len(str(header[c - 1]))
        ws.column_dimensions[letter].width = max(10, min(22, header_len + 2))
    ws.column_dimensions['A'].width = 32

    wb.save(path)
    return ws.title


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument('input', help='Excel de entrada (.xlsx) con stock/ventas por tienda')
    parser.add_argument('--sheet', default=None, help='Nombre de la hoja de entrada (por defecto, la primera)')
    parser.add_argument('--output', default=None, help='Excel de salida (por defecto, se escribe en el propio archivo de entrada)')
    parser.add_argument('--output-sheet', default='Traspasos', help='Nombre de la hoja nueva a crear/sobrescribir (por defecto "Traspasos")')
    parser.add_argument('--mapping', default=None, help='JSON con mapeo explícito de columnas, p.ej. {"sku": "Código", "stock_AMIGO": "Stock Amigo"}')
    parser.add_argument('--config', default=None, help='JSON con overrides de configuración de negocio (buffers, prefijos...)')
    args = parser.parse_args()

    cfg = dict(DEFAULT_CONFIG)
    if args.config:
        with open(args.config, encoding='utf-8') as f:
            cfg.update(json.load(f))

    explicit_mapping = {}
    if args.mapping:
        with open(args.mapping, encoding='utf-8') as f:
            explicit_mapping = json.load(f)

    sheet_title, headers, sheet_rows = read_input_sheet(args.input, args.sheet)
    mapping, missing_optional = detect_columns(headers, explicit_mapping)

    rows, stats = build_rows(sheet_rows, headers, mapping, cfg)
    if not rows:
        raise SystemExit('No se ha podido leer ninguna fila válida de datos (¿SKU vacío en todas las filas?).')

    compute_top20(rows, cfg)
    for r in rows:
        r['movimientos'] = compute_movimientos_para_fila(r, cfg)

    out_path = args.output or args.input
    if args.output and args.output != args.input:
        import shutil
        shutil.copyfile(args.input, args.output)
    written_sheet = write_output_sheet(out_path, args.output_sheet, rows, cfg)

    # ---- Resumen para el usuario ----
    n_accesorios = sum(1 for r in rows if is_accessory_row(r, cfg))
    n_con_traspaso = sum(1 for r in rows if any(m['qty'] > 0 for m in r['movimientos']))
    top20_counts = {s: sum(1 for r in rows if r['top20'][s]) for s in STORES}

    print(f'Hoja leída: "{sheet_title}" ({stats["filas_leidas"]} filas de datos).')
    print('Columnas detectadas:')
    for field, header in mapping.items():
        print(f'  - {field}: {header if header else "(no encontrada — se asume 0)"}')
    if missing_optional:
        print(f'Aviso: no se detectaron columnas opcionales para: {", ".join(missing_optional)} (se han tratado como 0).')
    if stats['wholesale']:
        print(f'Omitidas {stats["wholesale"]} filas Wholesale.')
    if stats['excluded']:
        print(f'Omitidas {stats["excluded"]} filas excluidas ({cfg["repoImportExcludeSkuPrefixes"]}).')
    if stats['negativos']:
        print(f'Aviso: {len(stats["negativos"])} valores negativos encontrados (revisa el origen): {stats["negativos"][:10]}{" ..." if len(stats["negativos"]) > 10 else ""}')
    print(f'{len(rows)} SKUs procesados ({n_accesorios} accesorios sin traspaso, {n_con_traspaso} con traspasos sugeridos).')
    print('Top20 por tienda: ' + ', '.join(f'{STORE_LABELS[s]}={top20_counts[s]}' for s in STORES))
    print(f'Hoja "{written_sheet}" escrita en: {out_path}')


if __name__ == '__main__':
    main()
